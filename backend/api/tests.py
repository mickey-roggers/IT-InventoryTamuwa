from io import BytesIO
from zipfile import ZipFile

from django.contrib.auth.models import User
from django.test import TestCase
from openpyxl import load_workbook
from rest_framework.test import APIClient

from assets.models import Asset, AssetLink, AssignmentHistory, Category, Department, Person, StatusOption
from issues.models import Issue, Project
from maintenance.models import MaintenanceLog, MaintenanceWorkEntry
from requisition.models import Requisition, RequisitionItem
from technicians.models import Technician
from users.models import UserProfile


class ExtendedApiTests(TestCase):
    def setUp(self):
        self.client = APIClient()
        self.department = Department.objects.create(name="ICT")
        self.admin = User.objects.create_user(
            username="api-admin", password="AdminPass123!", is_staff=False
        )
        UserProfile.objects.create(
            user=self.admin,
            role="admin",
            department=self.department,
            is_first_login=False,
        )
        self.viewer = User.objects.create_user(
            username="api-viewer", password="ViewerPass123!"
        )
        UserProfile.objects.create(
            user=self.viewer,
            role="viewer",
            is_first_login=False,
        )

    def test_user_management_is_readable_but_only_staff_can_write(self):
        self.client.force_authenticate(self.viewer)
        self.assertEqual(self.client.get("/api/users/").status_code, 200)
        denied = self.client.post(
            "/api/users/", {"username": "blocked", "password": "StrongPass123!"}, format="json"
        )
        self.assertEqual(denied.status_code, 403)

        self.client.force_authenticate(self.admin)
        created = self.client.post(
            "/api/users/",
            {
                "username": "new-user",
                "password": "StrongPass123!",
                "role": "technician",
                "department_id": self.department.pk,
                "is_active": True,
            },
            format="json",
        )
        self.assertEqual(created.status_code, 201, created.data)
        user = User.objects.get(username="new-user")
        self.assertEqual(user.profile.role, "technician")
        self.assertEqual(user.profile.department, self.department)
        self.assertTrue(user.profile.must_change_password)

    def test_current_user_can_update_profile_and_change_password(self):
        self.client.force_authenticate(self.viewer)
        updated = self.client.patch(
            "/api/auth/me/",
            {
                "first_name": "Api",
                "last_name": "Viewer",
                "phone_number": "+254700000000",
                "department_id": self.department.pk,
            },
            format="json",
        )
        self.assertEqual(updated.status_code, 200, updated.data)
        self.assertEqual(updated.data["profile"]["department_id"], self.department.pk)

        changed = self.client.post(
            "/api/auth/change-password/",
            {"current_password": "ViewerPass123!", "new_password": "AnotherStrongPass456!"},
            format="json",
        )
        self.assertEqual(changed.status_code, 200, changed.data)
        self.viewer.refresh_from_db()
        self.assertTrue(self.viewer.check_password("AnotherStrongPass456!"))

    def test_asset_quantities_report_status_totals(self):
        category = Category.objects.create(name="Laptop", short_code="LAP")
        available = StatusOption.objects.create(name="Available")
        in_use = StatusOption.objects.create(name="In Use")
        Asset.objects.create(
            asset_id="LAP-001", category=category, model_description="Laptop one",
            serial_number="SER-001", status=available, created_by=self.admin, updated_by=self.admin,
        )
        Asset.objects.create(
            asset_id="LAP-002", category=category, model_description="Laptop two",
            serial_number="SER-002", status=in_use, department=self.department,
            created_by=self.admin, updated_by=self.admin,
        )
        self.client.force_authenticate(self.viewer)
        response = self.client.get("/api/asset-quantities/")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data[0]["total"], 2)
        self.assertEqual(response.data[0]["available"], 1)
        self.assertEqual(response.data[0]["in_use"], 1)

    def test_asset_excel_export_returns_downloadable_workbook(self):
        category = Category.objects.create(name="Laptop", short_code="LAP")
        available = StatusOption.objects.create(name="Available")
        missing = StatusOption.objects.create(name="Missing")
        Asset.objects.create(
            asset_id="LAP-001",
            category=category,
            model_description="Export laptop",
            serial_number="EXPORT-SER-001",
            status=available,
            created_by=self.admin,
            updated_by=self.admin,
        )
        Asset.objects.create(
            asset_id="LAP-002",
            category=category,
            model_description="Missing assigned laptop",
            serial_number="EXPORT-SER-002",
            status=missing,
            department=self.department,
            created_by=self.admin,
            updated_by=self.admin,
        )
        self.client.force_authenticate(self.viewer)

        response = self.client.get("/api/assets/export/")

        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        self.assertEqual(
            response["Content-Disposition"],
            'attachment; filename="IT_Inventory_Export.xlsx"',
        )
        self.assertTrue(response.content.startswith(b"PK"))

        workbook = load_workbook(BytesIO(response.content), data_only=False)
        inventory = workbook["ASSET Inventory"]
        readme_text = "\n".join(
            str(cell.value or "") for cell in workbook["📋 README"]["A"]
        )

        self.assertEqual(inventory["H1"].value, "Current Status")
        self.assertEqual(inventory["K1"].value, "Status Override")
        asset_rows = {
            inventory.cell(row=row, column=1).value: row
            for row in range(2, 4)
        }
        available_row = asset_rows["LAP-001"]
        missing_row = asset_rows["LAP-002"]
        self.assertIn(
            f'IF(TRIM($F{available_row})<>"","In Use","Available")',
            inventory.cell(row=available_row, column=8).value,
        )
        self.assertIn(inventory.cell(row=available_row, column=11).value, (None, ""))
        self.assertEqual(inventory.cell(row=missing_row, column=11).value, "Missing")
        self.assertTrue(
            any(
                validation.formula1 == "LISTS!$B$2:$B$500"
                and "K2:K2000" in str(validation.sqref)
                for validation in inventory.data_validations.dataValidation
            )
        )
        self.assertEqual(workbook.calculation.calcMode, "auto")
        self.assertNotIn("ALT+F11", readme_text.upper())
        self.assertNotIn("PRIVATE SUB", readme_text.upper())
        self.assertNotIn("MACRO-ENABLED", readme_text.upper())

        with ZipFile(BytesIO(response.content)) as archive:
            self.assertNotIn("xl/vbaProject.bin", archive.namelist())

    def test_bought_queue_processing_links_assets_and_tracks_quantity(self):
        category = Category.objects.create(name="Monitor", short_code="MON")
        available = StatusOption.objects.create(name="Available")
        asset = Asset.objects.create(
            asset_id="MON-001", category=category, model_description="Monitor",
            serial_number="MON-SER-001", status=available,
            created_by=self.admin, updated_by=self.admin,
        )
        requisition = Requisition.objects.create(
            req_no="REQ-001", company="Tamuwa", title="Displays", status="Bought",
            created_by=self.admin,
        )
        item = RequisitionItem.objects.create(
            requisition=requisition, item_name="Monitor", item_type="Asset",
            unit_price="25000", quantity=2, is_approved=True,
        )
        self.client.force_authenticate(self.admin)
        queued = self.client.get("/api/requisition-items/?bought_queue=true")
        self.assertEqual(queued.data["count"], 1)

        first = self.client.post(
            f"/api/requisition-items/{item.pk}/process/", {"asset_id": asset.pk}, format="json"
        )
        self.assertEqual(first.status_code, 200, first.data)
        item.refresh_from_db()
        asset.refresh_from_db()
        self.assertEqual(item.quantity, 1)
        self.assertFalse(item.is_processed)
        self.assertEqual(asset.requisition, requisition)

        second = self.client.post(
            f"/api/requisition-items/{item.pk}/process/", {"asset_id": asset.pk}, format="json"
        )
        self.assertEqual(second.status_code, 200, second.data)
        item.refresh_from_db()
        self.assertTrue(item.is_processed)

    def test_technician_related_records_are_manageable(self):
        technician = Technician.objects.create(company_name="Repairs Ltd", technician_name="Jane")
        self.client.force_authenticate(self.viewer)
        denied = self.client.post(
            "/api/technician-assistants/",
            {"technician": technician.pk, "name": "Blocked", "is_active": True},
            format="json",
        )
        self.assertEqual(denied.status_code, 403)
        self.client.force_authenticate(self.admin)
        assistant = self.client.post(
            "/api/technician-assistants/",
            {"technician": technician.pk, "name": "Alex", "is_active": True},
            format="json",
        )
        self.assertEqual(assistant.status_code, 201, assistant.data)
        service = self.client.post(
            "/api/technician-services/",
            {"technician": technician.pk, "service_name": "Screen replacement", "typical_cost": "5000"},
            format="json",
        )
        self.assertEqual(service.status_code, 201, service.data)
        recommendation = self.client.post(
            "/api/technician-recommendations/",
            {
                "technician": technician.pk,
                "recommendation_type": "REPAIR",
                "description": "Replace damaged display",
                "priority": "HIGH",
            },
            format="json",
        )
        self.assertEqual(recommendation.status_code, 201, recommendation.data)

        detail = self.client.get(f"/api/technicians/{technician.pk}/")
        self.assertEqual(len(detail.data["assistants"]), 1)
        self.assertEqual(len(detail.data["services"]), 1)

    def test_asset_linking_accepts_multiple_assets_and_builds_direct_chain_edges(self):
        category = Category.objects.create(name="Dock", short_code="DCK")
        available = StatusOption.objects.create(name="Available")
        assets = [
            Asset.objects.create(
                asset_id=f"DCK-00{index}", category=category,
                model_description=f"Dock {index}", serial_number=f"DCK-SER-{index}",
                status=available, created_by=self.admin, updated_by=self.admin,
            )
            for index in range(1, 4)
        ]
        self.client.force_authenticate(self.admin)
        response = self.client.post(
            "/api/asset-links/",
            {"asset": assets[0].pk, "linked_asset_ids": [assets[1].pk, assets[2].pk]},
            format="json",
        )
        self.assertEqual(response.status_code, 201, response.data)
        self.assertEqual(AssetLink.objects.count(), 4)

    def test_creating_asset_from_bought_queue_generates_id_and_reduces_quantity(self):
        category = Category.objects.create(name="Laptop", short_code="LAP")
        available = StatusOption.objects.create(name="Available")
        requisition = Requisition.objects.create(
            req_no="REQ-QUEUE", company="Tamuwa", title="Laptops", status="Bought", created_by=self.admin
        )
        item = RequisitionItem.objects.create(
            requisition=requisition, item_type="Asset", item_name="Laptop",
            unit_price="50000", quantity=2, is_approved=True,
        )
        self.client.force_authenticate(self.admin)
        response = self.client.post(
            "/api/assets/",
            {
                "asset_id": "", "category_id": category.pk,
                "model_description": "Queue laptop", "serial_number": "QUEUE-SER-1",
                "status_id": available.pk, "requisition_item_id": item.pk,
            },
            format="json",
        )
        self.assertEqual(response.status_code, 201, response.data)
        self.assertEqual(response.data["asset_id"], "LAP-001")
        self.assertEqual(response.data["requisition"], requisition.pk)
        item.refresh_from_db()
        self.assertEqual(item.quantity, 1)
        self.assertFalse(item.is_processed)

    def test_asset_ownership_can_be_copied_to_the_full_linked_chain(self):
        category = Category.objects.create(name="Linked laptop", short_code="LLP")
        available = StatusOption.objects.create(name="Available")
        StatusOption.objects.create(name="In Use")
        person = Person.objects.create(first_name="Amina", last_name="Otieno", department=self.department)
        assets = [
            Asset.objects.create(
                asset_id=f"LLP-00{index}", category=category,
                model_description=f"Linked laptop {index}", serial_number=f"LLP-SER-{index}",
                status=available, created_by=self.admin, updated_by=self.admin,
            )
            for index in range(1, 3)
        ]
        AssetLink.objects.create(asset=assets[0], linked_asset=assets[1], created_by=self.admin)
        AssetLink.objects.create(asset=assets[1], linked_asset=assets[0], created_by=self.admin)
        self.client.force_authenticate(self.admin)

        response = self.client.patch(
            f"/api/assets/{assets[0].pk}/",
            {
                "assigned_to_id": person.pk,
                "department_id": self.department.pk,
                "sync_linked_ownership": True,
            },
            format="json",
        )

        self.assertEqual(response.status_code, 200, response.data)
        for asset in assets:
            asset.refresh_from_db()
            self.assertEqual(asset.assigned_to, person)
            self.assertEqual(asset.department, self.department)
            self.assertEqual(asset.status.name, "In Use")
            self.assertTrue(AssignmentHistory.objects.filter(asset=asset, person=person).exists())

    def test_requisition_creation_accepts_multiple_items(self):
        self.client.force_authenticate(self.admin)
        response = self.client.post(
            "/api/requisitions/",
            {
                "req_no": "REQ-MULTI",
                "company": "Tamuwa",
                "title": "New workstations",
                "status": "Pending",
                "items": [
                    {"item_type": "Asset", "item_name": "Laptop", "unit_price": "70000", "quantity": 2, "is_approved": True},
                    {"item_type": "Asset", "item_name": "Mouse", "unit_price": "1500", "quantity": 2, "is_approved": True},
                ],
            },
            format="json",
        )
        self.assertEqual(response.status_code, 201, response.data)
        self.assertEqual(Requisition.objects.get(pk=response.data["id"]).items.count(), 2)

    def test_maintenance_work_entries_keep_partial_repair_history(self):
        category = Category.objects.create(name="Repair device", short_code="RPD")
        available = StatusOption.objects.create(name="Available")
        asset = Asset.objects.create(
            asset_id="RPD-001", category=category, model_description="Repair device",
            serial_number="RPD-SER-1", status=available,
            created_by=self.admin, updated_by=self.admin,
        )
        maintenance = MaintenanceLog.objects.create(
            asset=asset, date_reported="2026-10-20", description="Intermittent boot issue",
            maintenance_status="Open", reported_by=self.admin,
        )
        self.client.force_authenticate(self.admin)
        response = self.client.post(
            "/api/maintenance-work-entries/",
            {"maintenance_log": maintenance.pk, "task": "Reseated memory", "outcome": "Booted once; monitoring continues."},
            format="json",
        )
        self.assertEqual(response.status_code, 201, response.data)
        entry = MaintenanceWorkEntry.objects.get(pk=response.data["id"])
        self.assertEqual(entry.recorded_by, self.admin)
        self.assertEqual(maintenance.work_entries.count(), 1)
        maintenance.refresh_from_db()
        self.assertEqual(maintenance.maintenance_status, "Open")

    def test_project_can_optionally_link_to_an_open_issue(self):
        issue = Issue.objects.create(title="Repeated network drops", reported_by=self.admin)
        self.client.force_authenticate(self.admin)
        response = self.client.post(
            "/api/projects/",
            {"title": "Replace core switch", "priority": "High", "status": "Pending", "linked_issue": issue.pk},
            format="json",
        )
        self.assertEqual(response.status_code, 201, response.data)
        project = Project.objects.get(pk=response.data["id"])
        self.assertEqual(project.linked_issue, issue)
